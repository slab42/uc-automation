#!/usr/bin/env python3
"""Build call_trees.xlsx from Cisco Unity call-handler and menu-entry CSV exports."""

from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

MAX_DEPTH = 50

CALL_HANDLER_FIELDS = {"displayName", "extension", "objectId", "uri"}
MENU_FIELDS = {
    "sourceHandler", "sourceExtension", "sourceObjectId", "touchToneKey",
    "destinationType", "destinationName", "targetHandlerObjectId",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FONT = Font(color="008000")
STATIC_FONT = Font(color="666666")
CAUTION_FILL = PatternFill("solid", fgColor="FCE4D6")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")


def prompt_path(prompt: str, default: str) -> Path:
    value = input(f"{prompt} ({default}): ").strip()
    return Path(value or default)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return [
            {str(k).strip(): (v or "").strip() for k, v in row.items()}
            for row in csv.DictReader(handle)
        ]


def require_fields(rows: list[dict[str, str]], required: set[str], label: str) -> None:
    if not rows:
        raise ValueError(f"{label} CSV contains no data rows")
    missing = required.difference(rows[0].keys())
    if missing:
        raise ValueError(f"{label} CSV is missing columns: {', '.join(sorted(missing))}")


def key_rank(value: str) -> tuple[int, str]:
    order = {str(i): i for i in range(10)}
    order.update({"*": 10, "#": 11})
    return order.get(value, 99), value


def is_routable(row: dict[str, str]) -> bool:
    """Keep assignments with an actual target, transfer, or conversation."""
    return bool(
        row.get("targetHandlerObjectId")
        or row.get("transferNumber")
        or row.get("targetConversation")
        or (
            row.get("destinationType")
            and row.get("destinationType") != "No configured destination"
        )
    )


def build_tree_rows(
    root_id: str,
    handlers_by_id: dict[str, dict[str, str]],
    edges_by_source: dict[str, list[dict[str, str]]],
    issues: list[dict[str, str]],
) -> list[dict[str, object]]:
    root = handlers_by_id[root_id]
    rows: list[dict[str, object]] = []

    def walk(current_id: str, level: int, path_ids: tuple[str, ...], path_text: str) -> None:
        current = handlers_by_id.get(current_id, {})
        current_name = current.get("displayName", current_id)

        if level > MAX_DEPTH:
            issues.append({
                "issueType": "Maximum depth exceeded",
                "sourceHandler": current_name,
                "sourceObjectId": current_id,
                "touchToneKey": "",
                "targetHandler": "",
                "targetObjectId": "",
                "details": f"Tree traversal exceeded MAX_DEPTH={MAX_DEPTH}",
            })
            return

        for edge in sorted(edges_by_source.get(current_id, []), key=lambda r: key_rank(r.get("touchToneKey", ""))):
            key = edge.get("touchToneKey", "")
            destination_type = edge.get("destinationType", "")
            target_id = edge.get("targetHandlerObjectId", "")
            destination = (
                edge.get("targetHandlerName")
                or edge.get("destinationName")
                or edge.get("transferDisplayName")
                or edge.get("transferNumber")
                or edge.get("targetConversation")
                or target_id
            )
            branch = f"Key {key} -> {destination}" if key else destination
            child_path = f"{path_text} -> {branch}"
            status = ""

            if target_id and target_id in path_ids:
                status = "CIRCULAR REFERENCE"
            elif (
                target_id and target_id not in handlers_by_id
                and destination_type != "Subscriber Mailbox"
            ):
                status = "TARGET NOT FOUND"

            rows.append({
                "rootHandler": root.get("displayName", root_id),
                "rootExtension": root.get("extension", ""),
                "level": level,
                "tree": f"{'    ' * level}|- {branch}",
                "path": child_path,
                "sourceHandler": current_name,
                "sourceObjectId": current_id,
                "touchToneKey": key,
                "destinationType": destination_type,
                "destinationName": destination,
                "transferNumber": edge.get("transferNumber", ""),
                "targetHandlerObjectId": target_id,
                "status": status,
            })

            if status:
                issues.append({
                    "issueType": status.title(),
                    "sourceHandler": current_name,
                    "sourceObjectId": current_id,
                    "touchToneKey": key,
                    "targetHandler": destination,
                    "targetObjectId": target_id,
                    "details": child_path,
                })
                continue

            if (
                target_id
                and destination_type == "Call Handler"
            ):
                walk(
                    target_id, level + 1, path_ids + (target_id,), child_path
                )

    rows.append({
        "rootHandler": root.get("displayName", root_id),
        "rootExtension": root.get("extension", ""),
        "level": 0,
        "tree": root.get("displayName", root_id),
        "path": root.get("displayName", root_id),
        "sourceHandler": root.get("displayName", root_id),
        "sourceObjectId": root_id,
        "touchToneKey": "",
        "destinationType": "Root Call Handler",
        "destinationName": root.get("displayName", root_id),
        "transferNumber": "",
        "targetHandlerObjectId": "",
        "status": "",
    })
    walk(root_id, 1, (root_id,), root.get("displayName", root_id))
    return rows


def write_sheet(ws, headers: list[str], rows: list[dict[str, object]], table_name: str) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"

    # Excel tables already provide their own filter controls.
    # Do not also assign ws.auto_filter.ref to the same range, because
    # Excel may repair/remove duplicate AutoFilter definitions.
    if rows:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)


def main() -> None:
    handler_csv = prompt_path("Call Handler CSV", "non_subscriber_callhandlers.csv")
    menu_csv = prompt_path("Menu Entry CSV", "menuentries.csv")
    output_xlsx = prompt_path("Excel Output File", "call_trees.xlsx")

    for path in (handler_csv, menu_csv):
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

    handlers = read_csv(handler_csv)
    menu_entries = read_csv(menu_csv)
    require_fields(handlers, CALL_HANDLER_FIELDS, "Call Handler")
    require_fields(menu_entries, MENU_FIELDS, "Menu Entry")

    handlers_by_id = {row["objectId"]: row for row in handlers if row.get("objectId")}
    routable_entries = [row for row in menu_entries if is_routable(row)]

    edges_by_source: dict[str, list[dict[str, str]]] = defaultdict(list)
    targeted_ids: set[str] = set()
    issues: list[dict[str, str]] = []

    for row in routable_entries:
        source_id = row.get("sourceObjectId", "")
        target_id = row.get("targetHandlerObjectId", "")
        if source_id:
            edges_by_source[source_id].append(row)
        if target_id:
            targeted_ids.add(target_id)
            if target_id not in handlers_by_id:
                issues.append({
                    "issueType": "Target not found",
                    "sourceHandler": row.get("sourceHandler", ""),
                    "sourceObjectId": source_id,
                    "touchToneKey": row.get("touchToneKey", ""),
                    "targetHandler": row.get("targetHandlerName", "") or row.get("destinationName", ""),
                    "targetObjectId": target_id,
                    "details": "Target ObjectId is not present in the Call Handlers CSV",
                })

    source_ids = set(edges_by_source)
    root_ids = sorted(
        source_ids - targeted_ids,
        key=lambda oid: handlers_by_id.get(oid, {}).get("displayName", oid).casefold(),
    )

    # Include routed components with no natural root, such as a fully cyclic component.
    covered: set[str] = set()
    tree_rows: list[dict[str, object]] = []
    for root_id in root_ids:
        component_rows = build_tree_rows(root_id, handlers_by_id, edges_by_source, issues)
        tree_rows.extend(component_rows)
        covered.update(
            str(row.get("sourceObjectId", ""))
            for row in component_rows
            if row.get("sourceObjectId")
        )

    for source_id in sorted(source_ids - covered):
        tree_rows.extend(build_tree_rows(source_id, handlers_by_id, edges_by_source, issues))

    roots = [
        {
            "displayName": handlers_by_id[oid].get("displayName", ""),
            "extension": handlers_by_id[oid].get("extension", ""),
            "objectId": oid,
            "outboundRoutes": len(edges_by_source.get(oid, [])),
        }
        for oid in root_ids
        if oid in handlers_by_id
    ]

    orphan_rows = [
        {
            "displayName": row.get("displayName", ""),
            "extension": row.get("extension", ""),
            "objectId": object_id,
            "reason": "Not a source and not targeted by another exported call handler",
        }
        for object_id, row in handlers_by_id.items()
        if object_id not in source_ids and object_id not in targeted_ids
    ]
    orphan_rows.sort(key=lambda row: row["displayName"].casefold())

    workbook = Workbook()
    workbook.remove(workbook.active)

    ws_summary = workbook.create_sheet("Summary")
    summary_rows = [
        ("Cisco Unity Call Tree Workbook", ""),
        ("Call handlers imported", len(handlers)),
        ("Menu-entry rows imported", len(menu_entries)),
        ("Routable menu entries", len(routable_entries)),
        ("Root handlers identified", len(root_ids)),
        ("Orphan handlers", len(orphan_rows)),
        ("Issues detected", len(issues)),
        ("Maximum traversal depth", MAX_DEPTH),
    ]
    for row in summary_rows:
        ws_summary.append(row)
    ws_summary["A1"].fill = TITLE_FILL
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary.column_dimensions["A"].width = 32
    ws_summary.column_dimensions["B"].width = 18
    for cell in ws_summary["B"]:
        cell.font = STATIC_FONT
    ws_summary.sheet_view.showGridLines = False

    ws_handlers = workbook.create_sheet("Call Handlers")
    write_sheet(ws_handlers, ["displayName", "extension", "objectId", "uri"], handlers, "CallHandlersTable")

    ws_menu = workbook.create_sheet("Menu Entries")
    menu_headers = list(menu_entries[0].keys()) if menu_entries else sorted(MENU_FIELDS)
    write_sheet(ws_menu, menu_headers, menu_entries, "MenuEntriesTable")

    ws_roots = workbook.create_sheet("Root Handlers")
    write_sheet(ws_roots, ["displayName", "extension", "objectId", "outboundRoutes"], roots, "RootHandlersTable")

    ws_tree = workbook.create_sheet("Call Trees")
    tree_headers = [
        "rootHandler", "rootExtension", "level", "tree", "path",
        "sourceHandler", "sourceObjectId", "touchToneKey", "destinationType",
        "destinationName", "transferNumber", "targetHandlerObjectId", "status",
    ]
    write_sheet(ws_tree, tree_headers, tree_rows, "CallTreesTable")
    ws_tree.column_dimensions["D"].width = 70
    ws_tree.column_dimensions["E"].width = 100
    for row in range(2, ws_tree.max_row + 1):
        level = ws_tree.cell(row, 3).value or 0
        ws_tree.cell(row, 4).alignment = Alignment(indent=min(int(level), 15))
        if ws_tree.cell(row, 13).value:
            for col in range(1, 14):
                ws_tree.cell(row, col).fill = ERROR_FILL

    ws_orphans = workbook.create_sheet("Orphan Handlers")
    write_sheet(ws_orphans, ["displayName", "extension", "objectId", "reason"], orphan_rows, "OrphanHandlersTable")

    ws_issues = workbook.create_sheet("Issues")
    issue_headers = [
        "issueType", "sourceHandler", "sourceObjectId", "touchToneKey",
        "targetHandler", "targetObjectId", "details",
    ]
    write_sheet(ws_issues, issue_headers, issues, "IssuesTable")
    for row in range(2, ws_issues.max_row + 1):
        for col in range(1, len(issue_headers) + 1):
            ws_issues.cell(row, col).fill = CAUTION_FILL

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_xlsx)

    print("\nWorkbook created successfully.")
    print(f"Output: {output_xlsx.resolve()}")
    print(f"Call handlers: {len(handlers)}")
    print(f"Routable menu entries: {len(routable_entries)}")
    print(f"Root handlers: {len(root_ids)}")
    print(f"Tree rows: {len(tree_rows)}")
    print(f"Issues: {len(issues)}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelled by user")
        sys.exit(130)
    except Exception as exc:
        print(f"\nERROR: {exc}")
        sys.exit(1)
